# ✅ Full updated code for Ask Mini Lawyer — 2025 Edition
import os
import streamlit as st
import torch
from openai import OpenAI
from dotenv import load_dotenv
from datetime import datetime
from app_resources import mongo_client, pinecone_client, model
import uuid
from streamlit_js import st_js, st_js_blocking
import json
import fitz
import docx
from fpdf import FPDF
import sys
import re

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

# Load environment variables
load_dotenv()
DATABASE_NAME = os.getenv("DATABASE_NAME")
client_openai = OpenAI(api_key=os.getenv("OPEN_AI"))

# External sources
judgment_index = pinecone_client.Index("judgments-names")
law_index = pinecone_client.Index("laws-names")
judgment_collection = mongo_client[DATABASE_NAME]["judgments"]
law_collection = mongo_client[DATABASE_NAME]["laws"]
conversation_collection = mongo_client[DATABASE_NAME]["conversations"]
document_feedback_collection = mongo_client[DATABASE_NAME]["document_feedback"]
chat_feedback_collection = mongo_client[DATABASE_NAME]["chat_feedback"]

torch.classes.__path__ = []
st.set_page_config(page_title="Ask Mini Lawyer", page_icon="💬", layout="wide")

# ===== UI Style =====
st.markdown("""
<style>
    .chat-container {background-color: #1E1E1E; padding: 20px; border-radius: 10px;}
    .chat-header {color: #4CAF50; font-size: 36px; font-weight: bold; text-align: center;}
    .user-message {background-color: #4CAF50; color: #ecf2f8; padding: 10px; border-radius: 10px; margin: 10px;}
    .bot-message {background-color: #44475a; color: #ecf2f8; padding: 10px; border-radius: 10px; margin: 10px;}
    .timestamp {font-size: 0.8em; color: #bbb;}
</style>
""", unsafe_allow_html=True)

# ===== Functions =====

def get_localstorage_value(key): return st_js_blocking(f"return localStorage.getItem('{key}');", key="get_" + key)
def set_localstorage_value(key, value): st_js(f"localStorage.setItem('{key}', '{value}');")

def get_or_create_chat_id():
    if 'current_chat_id' not in st.session_state:
        chat_id = get_localstorage_value("MiniLawyerChatId")
        if not chat_id or chat_id == "null":
            chat_id = str(uuid.uuid4())
            set_localstorage_value("MiniLawyerChatId", chat_id)
        st.session_state.current_chat_id = chat_id
    return st.session_state.current_chat_id

def save_conversation(chat_id, user_name, messages):
    conversation_collection.update_one(
        {"local_storage_id": chat_id},
        {"$set": {"local_storage_id": chat_id, "user_name": user_name, "messages": messages}},
        upsert=True
    )

def load_conversation(chat_id):
    convo = conversation_collection.find_one({"local_storage_id": chat_id})
    return convo.get('messages', []) if convo else []

def delete_conversation(chat_id):
    conversation_collection.delete_one({"local_storage_id": chat_id})
    st_js("localStorage.clear();")
    st.session_state.current_chat_id = None

def read_pdf(file):
    return "".join([page.get_text() for page in fitz.open(stream=file.read(), filetype="pdf")])

def read_docx(file):
    return "\n".join([p.text for p in docx.Document(file).paragraphs])

def show_typing_realtime(msg="🤖 הבוט מקליד..."):
    ph = st.empty()
    ph.markdown(f"<div style='color:gray;'>{msg}</div>", unsafe_allow_html=True)
    return ph

def add_message(role, content):
    st.session_state['messages'].append({
        "role": role, "content": content, "timestamp": datetime.now().strftime("%H:%M:%S")
    })

def save_document_feedback(chat_id, document_type, feedback):
    document_feedback_collection.insert_one({
        "chat_id": chat_id,
        "document_type": document_type,
        "feedback": feedback,
        "timestamp": datetime.now()
    })

def save_chat_feedback(chat_id, message_index, feedback):
    chat_feedback_collection.insert_one({
        "chat_id": chat_id,
        "message_index": message_index,
        "feedback": feedback,
        "timestamp": datetime.now()
    })

def detect_document_type(text):
    prompt = f"""האם המסמך הבא הוא חוזה, תביעה, החלטה או אחר? החזר תשובה קצרה בלבד:
---
{text[:3000]}
"""
    response = client_openai.chat.completions.create(
        model="gpt-4", messages=[{"role": "user", "content": prompt}], temperature=0
    )
    return response.choices[0].message.content.strip()

def split_into_sections(text):
    pattern = r'(סעיף\s+\d+|פרק\s+\d+|[\d]+\.\d+|[\d]+\))'
    parts = re.split(pattern, text)
    sections = []
    buffer = ""
    for part in parts:
        if re.match(pattern, part):
            if buffer:
                sections.append(buffer.strip())
            buffer = part
        else:
            buffer += " " + part
    if buffer:
        sections.append(buffer.strip())
    return [s for s in sections if len(s.strip()) > 30]

def find_relevant_judgments(text, top_k=3):
    try:
        embedding = model.encode([text], normalize_embeddings=True)[0]
        results = judgment_index.query(vector=embedding.tolist(), top_k=top_k, include_metadata=True)
        explanations = []
        for match in results["matches"]:
            meta = match.get("metadata", {})
            doc = judgment_collection.find_one({"CaseNumber": meta.get("CaseNumber")})
            if doc:
                name = doc.get("Name", "")
                desc = doc.get("Description", "")
                prompt = f"""סצנה:
{text}

פסק דין:
שם: {name}
תיאור: {desc}

מדוע פסק הדין רלוונטי לסיטואציה? דרג מ-0 עד 10 בפורמט JSON:
{{"advice": "הסבר", "score": 8}}"""
                reply = client_openai.chat.completions.create(
                    model="gpt-3.5-turbo", messages=[{"role": "user", "content": prompt}], temperature=0.5
                )
                parsed = json.loads(reply.choices[0].message.content.strip())
                explanations.append(f"פסק דין: {name}\nהסבר: {parsed['advice']} (ציון: {parsed['score']}/10)")
        return explanations
    except Exception as e:
        return [f"שגיאה באחזור פסקי דין: {e}"]

def find_relevant_laws(text, top_k=3):
    try:
        embedding = model.encode([text], normalize_embeddings=True)[0]
        results = law_index.query(vector=embedding.tolist(), top_k=top_k, include_metadata=True)
        explanations = []
        for match in results["matches"]:
            meta = match.get("metadata", {})
            doc = law_collection.find_one({"IsraelLawID": meta.get("IsraelLawID")})
            if doc:
                name = doc.get("Name", "")
                desc = doc.get("Description", "")
                prompt = f"""סצנה:
{text}

חוק:
שם: {name}
תיאור: {desc}

מדוע החוק רלוונטי לסיטואציה? החזר בפורמט JSON:
{{"advice": "הסבר", "score": 8}}"""
                reply = client_openai.chat.completions.create(
                    model="gpt-3.5-turbo", messages=[{"role": "user", "content": prompt}], temperature=0.5
                )
                parsed = json.loads(reply.choices[0].message.content.strip())
                explanations.append(f"חוק: {name}\nהסבר: {parsed['advice']} (ציון: {parsed['score']}/10)")
        return explanations
    except Exception as e:
        return [f"שגיאה באחזור חוקים: {e}"]

def export_pdf(filename="summary.pdf"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    if "detected_doc_type" in st.session_state:
        pdf.multi_cell(0, 10, f"סוג מסמך: {st.session_state['detected_doc_type']}\n\n")
    if "doc_summary" in st.session_state:
        pdf.multi_cell(0, 10, f"סיכום המסמך:\n{st.session_state['doc_summary']}\n\n")

    pdf.cell(0, 10, "שאלות ותשובות:", ln=True)
    for msg in st.session_state['messages']:
        role = "👤 שאלה" if msg['role'] == "user" else "🤖 תשובה"
        pdf.multi_cell(0, 10, f"{role}:\n{msg['content']}\n")

    if "doc_judgments" in st.session_state:
        pdf.cell(0, 10, "פסקי דין רלוונטיים:", ln=True)
        for j in st.session_state["doc_judgments"]:
            pdf.multi_cell(0, 10, f"- {j}")

    if "doc_laws" in st.session_state:
        pdf.cell(0, 10, "חוקים רלוונטיים:", ln=True)
        for l in st.session_state["doc_laws"]:
            pdf.multi_cell(0, 10, f"- {l}")

    filepath = os.path.join("outputs", filename)
    os.makedirs("outputs", exist_ok=True)
    pdf.output(filepath)
    return filepath

def display_messages():
    for i, msg in enumerate(st.session_state['messages']):
        role = "user-message" if msg['role'] == "user" else "bot-message"
        st.markdown(
            f"<div class='{role}'>{msg['content']}<div class='timestamp'>{msg['timestamp']}</div></div>",
            unsafe_allow_html=True
        )
        if msg["role"] == "assistant":
            col1, col2 = st.columns([1, 1])
            with col1:
                if st.button("👍", key=f"like_{i}"):
                    save_chat_feedback(st.session_state.current_chat_id, i, "like")
                    st.success("תודה על הדירוג!")
            with col2:
                if st.button("👎", key=f"dislike_{i}"):
                    save_chat_feedback(st.session_state.current_chat_id, i, "dislike")
                    st.warning("תודה, נשפר בהמשך.")

# ===== App Logic =====

st.markdown('<div class="chat-header">💬 Ask Mini Lawyer</div>', unsafe_allow_html=True)
chat_id = get_or_create_chat_id()

# Init session
if "user_name" not in st.session_state:
    st.session_state["user_name"] = None
if "messages" not in st.session_state:
    st.session_state["messages"] = load_conversation(chat_id)

# Login screen
if not st.session_state["user_name"]:
    with st.form("user_name_form"):
        name = st.text_input("הכנס שם להתחלת שיחה:")
        if st.form_submit_button("התחל שיחה") and name:
            st.session_state["user_name"] = name
            add_message("assistant", f"שלום {name}, איך אפשר לעזור?")
            save_conversation(chat_id, name, st.session_state["messages"])
            st.rerun()

else:
    with st.container():
        st.markdown('<div class="chat-container">', unsafe_allow_html=True)
        display_messages()
        st.markdown('</div>', unsafe_allow_html=True)

    uploaded_file = st.file_uploader("📄 העלה מסמך משפטי", type=["pdf", "docx"])
    if uploaded_file:
        st.session_state["uploaded_doc_text"] = read_pdf(uploaded_file) if uploaded_file.type == "application/pdf" else read_docx(uploaded_file)
        st.success("המסמך נטען בהצלחה!")

        with st.spinner("GPT מזהה את סוג המסמך..."):
            st.session_state["detected_doc_type"] = detect_document_type(st.session_state["uploaded_doc_text"])

        st.markdown(f"**סוג המסמך שהמערכת זיהתה:** `{st.session_state['detected_doc_type']}`")

        col1, col2 = st.columns([1, 1])
        with col1:
            if st.button("✅ נכון"):
                save_document_feedback(chat_id, st.session_state["detected_doc_type"], "correct")
                st.success("תודה על האישור!")
        with col2:
            if st.button("❌ שגוי"):
                save_document_feedback(chat_id, st.session_state["detected_doc_type"], "incorrect")
                st.warning("תודה על הדיווח, נשפר!")

    if "uploaded_doc_text" in st.session_state and st.button("📋 סכם את המסמך"):
        with st.spinner("GPT מסכם את המסמך..."):
            summary_prompt = f"""סכם את המסמך המשפטי הבא בקצרה:
---
{st.session_state['uploaded_doc_text']}
"""
            response = client_openai.chat.completions.create(
                model="gpt-4", messages=[{"role": "user", "content": summary_prompt}], temperature=0.5
            )
            st.session_state["doc_summary"] = response.choices[0].message.content.strip()

    if "doc_summary" in st.session_state:
        st.markdown("### סיכום המסמך:")
        st.info(st.session_state["doc_summary"])

        if st.button("📚 הצג חוקים ופסקי דין למסמך"):
            with st.spinner("מאחזר פסקי דין וחוקים למסמך..."):
                st.session_state["doc_judgments"] = find_relevant_judgments(st.session_state["doc_summary"])
                st.session_state["doc_laws"] = find_relevant_laws(st.session_state["doc_summary"])

            st.subheader("📚 פסקי דין שנמצאו:")
            for j in st.session_state.get("doc_judgments", []):
                st.markdown(f"- {j}")
            st.subheader("⚖️ חוקים שנמצאו:")
            for l in st.session_state.get("doc_laws", []):
                st.markdown(f"- {l}")

        if st.button("🔍 ניתוח לפי סעיפים"):
            sections = split_into_sections(st.session_state["uploaded_doc_text"])
            for i, sec in enumerate(sections):
                st.markdown(f"#### סעיף {i+1}: {sec[:100]}...")
                with st.expander("הצג סעיף"):
                    st.write(sec)

        if st.button("📄 ייצא הכל כ-PDF"):
            path = export_pdf()
            with open(path, "rb") as f:
                st.download_button("📅 הורד PDF", f, file_name="legal_summary.pdf")

    # Chat with the bot
    with st.form("chat_form"):
        user_input = st.text_area("הכנס שאלה משפטית", height=100)
        if st.form_submit_button("שלח שאלה") and user_input.strip():
            add_message("user", user_input)
            save_conversation(chat_id, st.session_state["user_name"], st.session_state["messages"])
            st.rerun()

    if st.session_state['messages'] and st.session_state['messages'][-1]['role'] == "user":
        typing = show_typing_realtime()
        response = client_openai.chat.completions.create(
            model="gpt-4",
            messages=[{"role": "system", "content": "אתה עוזר משפטי מקצועי בדין הישראלי. ענה בקצרה ומדויק."}] +
                     [{"role": m["role"], "content": m["content"]} for m in st.session_state["messages"][-5:]] +
                     [{"role": "user", "content": st.session_state['messages'][-1]['content']}],
            max_tokens=700,
            temperature=0.7
        )
        typing.empty()
        add_message("assistant", response.choices[0].message.content.strip())
        save_conversation(chat_id, st.session_state["user_name"], st.session_state["messages"])
        st.rerun()

    if st.button("🗑 נקה שיחה"):
        delete_conversation(chat_id)
        st.session_state["messages"] = []
        st.session_state["user_name"] = None
        st.rerun()
